"""CSV/XLSX 结构化数据导入与确定性关联查询（不调用 LLM/Embedding）。"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from collections import defaultdict
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from django.db import transaction
from django.db.models import Q

from .models import Document, StructuredDataset, StructuredRecord


FIELD_LABELS = {
    "drawing_no": "图纸号",
    "part_no": "部件号",
    "part_name": "部件名称",
    "g_code": "G-code",
    "scp_level": "SCP 等级",
    "equipment": "设备/资产",
    "apex_no": "APEX 编号",
}

HEADER_ALIASES = {
    "drawing_no": [
        "图纸号", "图号", "图纸编号", "图样编号", "drawing no", "drawing number",
        "drawing_no", "drawing", "dwg no", "dwg number", "dwg",
    ],
    "part_no": [
        "部件号", "部件编号", "零件号", "零件编号", "物料号", "物料编码",
        "part no", "part number", "part_no", "component no", "component number",
    ],
    "part_name": [
        "部件名称", "零件名称", "物料名称", "名称", "part name", "part_name",
        "component name", "component", "part description", "description",
    ],
    "g_code": ["g-code", "g code", "gcode", "g_code", "g编码", "g代码"],
    "scp_level": ["scp等级", "scp级别", "scp level", "scp_level", "scp"],
    "equipment": [
        "设备", "设备号", "设备编号", "设备名称", "资产", "资产号", "资产编号",
        "equipment", "equipment no", "asset", "asset no", "machine", "facility", "attraction",
    ],
    "apex_no": ["apex", "apex号", "apex编号", "apex id", "apex no", "apex number"],
}


class StructuredDataError(ValueError):
    pass


def value_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def normalize_key(value: Any) -> str:
    """规范化业务键：忽略空格、横线等排版差异，统一大写。"""
    text = value_to_text(value)
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "", text).upper()


def _normalize_header(value: Any) -> str:
    return normalize_key(value).lower()


_ALIAS_LOOKUP = {
    _normalize_header(alias): canonical
    for canonical, aliases in HEADER_ALIASES.items()
    for alias in aliases
}


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return value_to_text(value)


def _unique_headers(values: Iterable[Any]) -> list[str]:
    headers: list[str] = []
    used: dict[str, int] = {}
    for index, value in enumerate(values, 1):
        base = value_to_text(value) or f"未命名列{index}"
        used[base] = used.get(base, 0) + 1
        headers.append(base if used[base] == 1 else f"{base}_{used[base]}")
    return headers


def detect_mapping(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for header in headers:
        canonical = _ALIAS_LOOKUP.get(_normalize_header(header))
        if canonical and canonical not in mapping:
            mapping[canonical] = header
    return mapping


def _parse_csv(payload: bytes) -> list[tuple[str, int, dict[str, Any], dict[str, str]]]:
    decoded = None
    for encoding in ("utf-8-sig", "gb18030", "utf-16", "latin-1"):
        try:
            decoded = payload.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise StructuredDataError("无法识别 CSV 编码，请保存为 UTF-8 后重试。")

    sample = decoded[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(decoded), dialect))
    return _rows_from_matrix(rows, "CSV")


def _parse_xlsx(payload: bytes) -> list[tuple[str, int, dict[str, Any], dict[str, str]]]:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise StructuredDataError(f"无法读取 Excel：{exc}") from exc

    parsed: list[tuple[str, int, dict[str, Any], dict[str, str]]] = []
    try:
        for sheet in workbook.worksheets:
            parsed.extend(_rows_from_matrix(sheet.iter_rows(values_only=True), sheet.title))
    finally:
        workbook.close()
    return parsed


def _rows_from_matrix(
    matrix: Iterable[Iterable[Any]], sheet_name: str,
) -> list[tuple[str, int, dict[str, Any], dict[str, str]]]:
    iterator = iter(matrix)
    # 业务 Excel 常在表头前放标题/说明。扫描前 20 行，优先选择能识别出最多
    # 标准关联字段的一行；完全无法识别时才回退到第一条非空行。
    buffered: list[tuple[int, list[Any]]] = []
    for row_number, row in enumerate(iterator, 1):
        buffered.append((row_number, list(row)))
        if row_number >= 20:
            break
    nonempty = [item for item in buffered if any(value_to_text(v) for v in item[1])]
    if not nonempty:
        return []

    scored = []
    for row_number, values in nonempty:
        candidate_headers = _unique_headers(values)
        scored.append((len(detect_mapping(candidate_headers)), -row_number, row_number, values))
    _score, _neg_row, header_row_number, header_values = max(scored)

    headers = _unique_headers(header_values)
    mapping = detect_mapping(headers)
    output = []

    remaining_rows = [item for item in buffered if item[0] > header_row_number]
    remaining_rows.extend((row_number, list(row)) for row_number, row in enumerate(iterator, len(buffered) + 1))
    for row_number, values in remaining_rows:
        values += [None] * max(0, len(headers) - len(values))
        data = {headers[i]: _json_value(values[i]) for i in range(len(headers))}
        if not any(value_to_text(v) for v in data.values()):
            continue
        output.append((sheet_name, row_number, data, mapping))
    return output


def parse_upload(upload) -> tuple[bytes, list[tuple[str, int, dict[str, Any], dict[str, str]]]]:
    payload = upload.read()
    if not payload:
        raise StructuredDataError("上传文件为空。")
    if len(payload) > 25 * 1024 * 1024:
        raise StructuredDataError("文件超过 25 MB，请拆分后再导入。")
    ext = Path(upload.name).suffix.lower()
    if ext in (".csv", ".tsv"):
        return payload, _parse_csv(payload)
    if ext in (".xlsx", ".xlsm"):
        return payload, _parse_xlsx(payload)
    if ext == ".xls":
        raise StructuredDataError("暂不支持旧版 .xls，请在 Excel 中另存为 .xlsx。")
    raise StructuredDataError("仅支持 CSV、TSV、XLSX 和 XLSM 文件。")


def _mapped_value(data: dict[str, Any], mapping: dict[str, str], canonical: str) -> str:
    header = mapping.get(canonical)
    return value_to_text(data.get(header)) if header else ""


@transaction.atomic
def import_structured_dataset(upload, kind: str, user=None) -> StructuredDataset:
    valid_kinds = {value for value, _label in StructuredDataset.Kind.choices}
    if kind not in valid_kinds:
        raise StructuredDataError("请选择正确的数据类型。")

    payload, rows = parse_upload(upload)
    if not rows:
        raise StructuredDataError("文件中没有可导入的数据行。")

    source_name = Path(upload.name).name[:255]
    checksum = hashlib.sha256(payload).hexdigest()
    duplicate = StructuredDataset.objects.filter(
        kind=kind, checksum=checksum, active=True,
    ).first()
    if duplicate:
        raise StructuredDataError(f"相同文件已经导入：{duplicate.source_name}")

    StructuredDataset.objects.filter(
        kind=kind, source_name=source_name, active=True,
    ).update(active=False)

    combined_mapping: dict[str, list[str]] = defaultdict(list)
    for _sheet, _row_number, _data, mapping in rows:
        for canonical, header in mapping.items():
            if header not in combined_mapping[canonical]:
                combined_mapping[canonical].append(header)

    dataset = StructuredDataset.objects.create(
        name=Path(source_name).stem[:160],
        kind=kind,
        source_name=source_name,
        checksum=checksum,
        row_count=len(rows),
        mapping=dict(combined_mapping),
        imported_by=user if getattr(user, "is_authenticated", False) else None,
    )

    records = []
    for sheet_name, row_number, data, mapping in rows:
        values = {field: _mapped_value(data, mapping, field) for field in FIELD_LABELS}
        records.append(StructuredRecord(
            dataset=dataset,
            sheet_name=sheet_name,
            row_number=row_number,
            drawing_no=values["drawing_no"],
            drawing_no_norm=normalize_key(values["drawing_no"]),
            part_no=values["part_no"],
            part_no_norm=normalize_key(values["part_no"]),
            part_name=values["part_name"],
            part_name_norm=normalize_key(values["part_name"]),
            g_code=values["g_code"],
            g_code_norm=normalize_key(values["g_code"]),
            scp_level=values["scp_level"],
            equipment=values["equipment"],
            equipment_norm=normalize_key(values["equipment"]),
            apex_no=values["apex_no"],
            apex_no_norm=normalize_key(values["apex_no"]),
            raw_data=data,
        ))
    StructuredRecord.objects.bulk_create(records, batch_size=1000)
    return dataset


def _relation_query(keys: dict[str, set[str]]) -> Q:
    query = Q(pk__in=[])
    for field in ("drawing_no", "part_no", "part_name", "g_code", "equipment", "apex_no"):
        values = keys.get(field) or set()
        if values:
            query |= Q(**{f"{field}_norm__in": values})
    return query


def _direct_lookup_query(normalized: str, partial: bool = False) -> Q:
    """在所有标准业务键中查找用户输入，而不只局限于图纸号。"""
    query = Q(pk__in=[])
    lookup = "icontains" if partial else "exact"
    for field in ("drawing_no", "part_no", "part_name", "g_code", "equipment", "apex_no"):
        query |= Q(**{f"{field}_norm__{lookup}": normalized})
    return query


def lookup_related(query_text: str, department: str = "") -> dict[str, Any]:
    """统一搜索业务键，并沿关联键进行最多三轮确定性跨表关联。

    department：手册全文搜索按部门过滤（空 = 不过滤，仅内部/兼容用途；
    页面调用时传当前用户部门，可见 通用 ∪ 该部门 的文档）。
    """
    normalized = normalize_key(query_text)
    if not normalized:
        return {
            "records": [], "groups": [], "manual_matches": [],
            "related_keys": {}, "match_mode": "none", "direct_count": 0,
        }

    base = StructuredRecord.objects.filter(dataset__active=True).select_related("dataset")
    records = list(base.filter(_direct_lookup_query(normalized))[:500])
    match_mode = "exact"
    if not records:
        records = list(base.filter(_direct_lookup_query(normalized, partial=True))[:500])
        match_mode = "partial"
    if not records:
        match_mode = "none"
    direct_count = len(records)

    keys: dict[str, set[str]] = {
        "drawing_no": set(), "part_no": set(), "part_name": set(),
        "g_code": set(), "equipment": set(), "apex_no": set(),
    }
    seen_ids = {record.id for record in records}
    for _round in range(3):
        for record in records:
            for field in keys:
                value = getattr(record, f"{field}_norm", "")
                if value:
                    keys[field].add(value)
        related = list(base.filter(_relation_query(keys)).exclude(id__in=seen_ids)[:1000])
        if not related:
            break
        records.extend(related)
        seen_ids.update(record.id for record in related)

    groups_by_kind: dict[str, list[StructuredRecord]] = defaultdict(list)
    for record in records:
        groups_by_kind[record.dataset.kind].append(record)
    kind_labels = dict(StructuredDataset.Kind.choices)
    groups = [
        {
            "kind": kind,
            "label": kind_labels.get(kind, kind),
            "records": rows[:100],
            "total": len(rows),
            "truncated": len(rows) > 100,
        }
        for kind, _label in StructuredDataset.Kind.choices
        if (rows := groups_by_kind.get(kind))
    ]

    doc_qs = Document.objects.filter(status=Document.Status.COMPLETED)
    if department:
        from .models import DEPARTMENT_GENERAL
        doc_qs = doc_qs.filter(kb__department__in=[DEPARTMENT_GENERAL, department])
    docs = list(
        doc_qs.select_related("kb").only("id", "original_name", "md_content", "kb__name")
    )
    from .inspection import search_in_content
    manual_matches = search_in_content(query_text, docs, limit=30)

    display_keys = {
        FIELD_LABELS[field]: sorted(values)
        for field, values in keys.items() if values
    }
    return {
        "records": records,
        "groups": groups,
        "manual_matches": manual_matches,
        "related_keys": display_keys,
        "match_mode": match_mode,
        "direct_count": direct_count,
    }


def lookup_drawing(query_text: str) -> dict[str, Any]:
    """向后兼容旧调用；现在实际执行统一关联搜索。"""
    return lookup_related(query_text)
